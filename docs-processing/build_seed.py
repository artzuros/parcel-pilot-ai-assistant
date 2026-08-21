import json, pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / 'data' / 'source' / 'ParcelPilot_Assessment_Data.xlsx'
OUT = ROOT / 'data' / 'seed'

SUSPECT_RESOLUTIONS = {"TKT-450", "TKT-451"}
INT_FIELDS = {"shipment_fee_inr"}

def read_sheet(ws):
    headers = [str(c.value).strip() if c.value is not None 
               else None for c in next(ws.iter_rows())]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row): continue
        rec = {}
        for header, value in zip(headers,row):
            if header is None: continue
            if value is None: rec[header] = None
            elif header in INT_FIELDS and isinstance(value, float):
                rec[header] = int(value)
            elif isinstance(value, str):
                rec[header] = value.strip()
            else: rec[header] = value
        rows.append(rec)
    return rows

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    
    for name in ("accounts", "orders", "tickets"):
        if name not in wb.sheetnames:
            raise SystemExit(f"Missing Sheet {name!r} in {SRC.name}")
        rows = read_sheet(wb[name])
        out = OUT / f"{name}.json"
        out.write_text(json.dumps(rows, indent=2, ensure_ascii=False),
                       encoding='utf-8')
        print(f"{name}: {len(rows)} rows TO {out.name}")
    
    tickets = json.loads((OUT / "tickets.json").read_text(encoding="utf-8"))
    for t in tickets:
        if t.get("historical_resolution") and t.get("ticket_id") in SUSPECT_RESOLUTIONS:
            t["resolution_confidence"] = "suspect"
        elif t.get("historical_resolution"):
            t["resolution_confidence"] = "ok"
        else: t["resolution_confidence"] = None
    (OUT / "tickets.json").write_text(
        json.dumps(tickets, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    wb_readme = read_sheet(wb["README"])
    meta = {r["Dataset snapshot"] if "Dataset snapshot" in r
            else k: v for r in wb_readme for k,v in r.items()} if wb_readme else {} 
    meta = {"snapshot": "2026-08-16 11:00 Asia/Kolkata",
            "currency": "INR",
            "notes": "Synthetic dataset created for a hiring assessment."}
    (OUT / "meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print(f"meta: {meta} -> meta.json")
    

if __name__=="__main__": main()